#!/usr/bin/env python3
"""Build docs/exercise-template.xlsx: every input a new exercise needs, with the
glute bridge filled in as the example. Run: python3 scripts/exercise-template.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

FONT = 'Arial'
wb = Workbook()
ink = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
title = Font(name=FONT, size=14, bold=True)
blue = Font(name=FONT, size=10, color='0000FF')          # example values: what a filled row looks like
muted = Font(name=FONT, size=9, color='666666', italic=True)
fill_head = PatternFill('solid', fgColor='1F3A5F')
fill_in = PatternFill('solid', fgColor='FFF7CC')          # yellow: fill this in
fill_ex = PatternFill('solid', fgColor='EEF3FA')
thin = Side(style='thin', color='BBBBBB')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

def sheet(name, widths):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws

def header_row(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c); cell.font = head; cell.fill = fill_head; cell.alignment = wrap; cell.border = box
    ws.row_dimensions[row].height = 30

def table(ws, start, cols, example_rows, blank_rows, note=None):
    header_row(ws, start, cols)
    r = start + 1
    for ex in example_rows:
        for i, v in enumerate(ex, 1):
            cell = ws.cell(row=r, column=i, value=v); cell.font = blue; cell.fill = fill_ex; cell.alignment = wrap; cell.border = box
        r += 1
    for _ in range(blank_rows):
        for i in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=i); cell.fill = fill_in; cell.border = box; cell.alignment = wrap; cell.font = ink
        r += 1
    if note:
        c = ws.cell(row=r + 1, column=1, value=note); c.font = muted; c.alignment = wrap
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=min(len(cols), 8))
    return r

def dv(ws, formula, rng, prompt=None):
    d = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
    if prompt: d.prompt = prompt; d.showInputMessage = True
    ws.add_data_validation(d); d.add(rng)

# ---------------------------------------------------------------- Read me
ws = wb.active; ws.title = 'Read me'
ws.column_dimensions['A'].width = 110; ws.sheet_view.showGridLines = False
lines = [
    ('OnTrack — setting up a new exercise', title),
    ('One sheet per part of an exercise. Fill the yellow cells; the blue rows are the glute bridge, filled in as it runs in the app today, so every column has a real example. Nothing on these sheets is read by the app directly: they are the brief a move is written from (public/js/moves.js), and the checklist at the end is what has to be true before it is switched on.', ink),
    ('', ink),
    ('Sheets', bold),
    ('1. Exercise — identity, the phone, the set (hold or reps, sets, targets), the opening words, the starting position and the rule that says the person is in it, what a rep is.', ink),
    ('2. Measurements — every angle or height the camera reads, which landmarks it is taken from, its allowed band, and how it is shown.', ink),
    ('3. Faults — every correction the coach can make: the measurement it comes from, which side of the band, the words (short, spoken, stronger), whether it is coached before the movement starts, and its order — the earliest in the chain of cause is said first.', ink),
    ('4. Words & sounds — the fixed cues (hold, lower, count, done, can\'t see you) and the tone each kind of cue gets.', ink),
    ('5. Muscles — what works, and how hard, for the muscle figure.', ink),
    ('6. Figure — the stick figure\'s two keyframes as joint angles, so the demo animation can be drawn.', ink),
    ('7. Checklist — recordings, tuning on the Review page, and sign-off.', ink),
    ('', ink),
    ('How the coach uses these', bold),
    ('Opening: the start text is said once. Nothing else is said until the person has held the starting position for the set-up wait (2 s by default, seen the whole time). If nobody is in the frame, "I can\'t see you — step into the camera" is said, and again every 15 s while that stays so.', ink),
    ('Coaching: every measurement is read on every frame and smoothed. A fault is a measurement outside its band; one is said only after it has held for the persist time (0.5 s), the same one not again inside the cooldown (4 s), and no two inside the gap (1.5 s). When several are wrong, the one earliest in the Faults order is said, whatever their sizes. Set-up faults are coached at the start position, before the movement is asked for. Words shown on the picture: every fault present, whether or not it is the one being said. Only corrections are red.', ink),
    ('Reps: from the start position the movement is asked for (the prompt); "raised" is when the progress measurement passes the raise threshold; the hold at the top runs once every measurement in "in position" is good; "Lower slowly" once the hold is done; the rep counts on the return to the start (below the down threshold). Back down before the hold is done: "Hold it at the top next time", not counted. A lowering faster than the slow-lowering seconds gets "slower on the way down" with the count.', ink),
    ('Holds: the clock runs while every measurement in "in position" is good, after a short settle; it stops the instant one is not. The time is called at the listed marks; the target ends the set.', ink),
    ('', ink),
    ('Legend', bold),
    ('Yellow cells: yours to fill.   Blue text: the example.   A dropdown arrow: pick from the list (the app only knows those).', ink),
]
for i, (text, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=text); c.font = f; c.alignment = wrap
    if f is title: ws.row_dimensions[i].height = 24

# ---------------------------------------------------------------- Exercise
ws = sheet('Exercise', [34, 46, 60, 46])
header_row(ws, 1, ['Field', 'Your answer', 'What it is for', 'Example — glute bridge'])
rows = [
    ('IDENTITY', None, None, None),
    ('id', 'One word, lowercase, no spaces: the key in the code and in the settings store.', 'bridge'),
    ('Name', 'As shown on the page and said by the coach.', 'Glute bridge'),
    ('One-line description', 'For the exercise card.', 'Lying on your back, hips lifted until knees, hips and shoulders are in a line.'),
    ('Category', 'Where it sits in a list.', 'Glutes and hips'),
    ('Equipment', 'Anything needed beyond the floor and a wall.', 'None (a mat helps)'),
    ('THE PHONE', None, None, None),
    ('Phone orientation', 'How the phone lies. The picture takes the whole screen once the phone is that way round. Pick one.', 'On its side (wide)'),
    ('Phone placement, in words', 'Said and shown before the set. Where, how far, at what height.', 'Lay the phone on its side on the floor, two or three metres away.'),
    ('View', 'Which way the body faces the camera. Every measurement below must be readable from this view.', 'Side on'),
    ('Which side faces the camera', 'Either, nearest limb, left, right. The app follows the side it sees best.', 'Either — the app follows the nearer leg'),
    ('THE SET', None, None, None),
    ('Type', 'A hold (one position, timed) or reps (a movement, counted). Pick one.', 'Reps'),
    ('Hold seconds', 'A hold: the target for the set. Reps: the hold at the top of each rep.', '2'),
    ('Reps per set', 'Reps only.', '10'),
    ('Sets', 'Default number of sets in a session.', '3'),
    ('Alternate sides between sets', 'Yes for one-leg or one-arm moves done a set per side.', 'No'),
    ('Slow lowering, seconds', 'Reps only: a return quicker than this is remarked on with the count. Blank for no rule.', '1'),
    ('Rest after a rep, seconds', 'The quiet after a count before the next rep is asked for.', '2'),
    ('Time calls', 'Holds: seconds left at which the time is called out.', '45, 30, 10, 5'),
    ('THE OPENING', None, None, None),
    ('Start text (what the coach says first)', 'Under 120 characters. Says where the phone goes and what position to get into. Nothing else is said until the starting position is held.', 'Lay the phone on its side on the floor. I will wait while you get set up: lie down side on to it, knees bent.'),
    ('Starting position, in words', 'Shown on the set-up card. What the body looks like before the movement.', 'Lie on your back, side on to the phone, knees bent, feet flat on the floor, arms by your sides.'),
    ('Starting position, as a rule', 'What the camera checks to say the person is in it: which measurements, which values. Held for the set-up wait before coaching begins.', 'Hips down (hip angle under the down threshold) and the shin standing up off the heel (shin 45–150°).'),
    ('Set-up wait, seconds', 'How long the starting position is held before coaching begins.', '2'),
    ('THE MOVEMENT (reps only)', None, None, None),
    ('Progress measurement', 'The one measurement that says how far into the rep the person is.', 'hip (knee–hip–shoulder angle)'),
    ('Raise threshold', 'Progress value past which the rep is under way (the top is being reached).', '150'),
    ('Down threshold', 'Progress value under which the person is back at the start; the rep counts here.', '135'),
    ('In position means', 'Which measurements must all be good for the hold at the top (or the hold, for a hold move) to run.', 'shin, hip, over, foot'),
    ('What counts a rep', 'In words, for the reader.', 'Hips lifted to the line, held two seconds, lowered slowly, back on the floor.'),
    ('SAFETY AND CONTEXT', None, None, None),
    ('Safety notes', 'What to stop for; what the coach cannot see.', 'Stop if the lower back pinches. The camera cannot see the spine between hip and shoulder.'),
    ('Common mistakes (for the reader)', 'The faults, in plain words, for the About text.', 'Pushing the hips past the knees; heels lifting; feet too far out.'),
    ('Easier version', '', 'Smaller lift, shorter hold.'),
    ('Harder version', '', 'Single leg; a longer hold at the top.'),
]
r = 2
for row in rows:
    if row[1] is None:
        c = ws.cell(row=r, column=1, value=row[0]); c.font = head; c.fill = fill_head
        for i in range(2, 5): ws.cell(row=r, column=i).fill = fill_head
        r += 1; continue
    field, why, ex = row
    ws.cell(row=r, column=1, value=field).font = bold
    c = ws.cell(row=r, column=2); c.fill = fill_in; c.border = box; c.alignment = wrap; c.font = ink
    ws.cell(row=r, column=3, value=why).font = ink; ws.cell(row=r, column=3).alignment = wrap
    e = ws.cell(row=r, column=4, value=ex); e.font = blue; e.alignment = wrap
    if field == 'Phone orientation': dv(ws, '"On its side (wide),Stood up (tall)"', f'B{r}')
    if field == 'View': dv(ws, '"Side on,Front on,Three-quarter"', f'B{r}')
    if field == 'Type': dv(ws, '"Hold,Reps"', f'B{r}')
    if field == 'Alternate sides between sets': dv(ws, '"Yes,No"', f'B{r}')
    r += 1
ws.freeze_panes = 'A2'

# ---------------------------------------------------------------- Measurements
ws = sheet('Measurements', [12, 20, 40, 22, 14, 14, 14, 10, 10, 10, 10, 16, 16, 18, 34])
cols = ['key', 'Name on screen', 'What it measures, in words', 'Kind', 'Landmark A', 'Landmark B (the joint)', 'Landmark C', 'Band low', 'Band high', 'Scale low', 'Scale high', 'Judged when', 'Part of "in position"?', 'Drawn as', 'Notes (why this band; what a wrong reading looks like)']
ex = [
    ['shin', 'SHIN', 'The shin\'s angle at the heel, between the toe and the knee: where the feet are.', 'Angle at joint B between A and C', 'toe', 'heel', 'knee', 85, 110, 50, 130, 'Start position (set-up)', 'Yes', 'Arc with number', 'Over the band the feet are out too far; under it, too close. Not judged while a heel or toe is off the floor.'],
    ['hip', 'HIP', 'The angle at the hip between knee and shoulder: the lift.', 'Angle at joint B between A and C', 'knee', 'hip', 'shoulder', 160, 180, 90, 180, 'At the top', 'Yes', 'Guide line knee–shoulder', 'The progress measurement for the rep.'],
    ['over', 'HEIGHT', 'How far the hip rises above the knee, as degrees of the thigh line off level.', 'Rise of C over A (signed)', 'knee', '', 'hip', -40, 3, -40, 20, 'At the top', 'Yes', 'Floor line through the knee', 'Above +3 the hips are past the knees (lower back arching).'],
    ['foot', 'FOOT', 'The foot\'s line off the floor, heel to toe.', 'Angle from the floor (signed)', 'heel', '', 'toe', -10, 10, -30, 30, 'Always', 'Yes', 'Readout beside the foot', 'Positive: heel up. Negative: toes up. Checked before the shin.'],
]
last = table(ws, 1, cols, ex, 10, 'Kinds the app knows: angle at a joint (three landmarks), tilt from vertical (two), angle from the floor (two), rise of one point over another (two, signed), bend of a three-point line off straight. Landmarks: nose, ear, shoulder, elbow, wrist, hip, knee, ankle, heel, toe — the side facing the camera is chosen for you.')
dv(ws, '"Angle at joint B between A and C,Tilt from vertical (A over B),Angle from the floor (signed),Rise of C over A (signed),Bend of line A–B–C off straight"', f'D2:D{last}')
for col in 'EFG': dv(ws, '"nose,ear,shoulder,elbow,wrist,hip,knee,ankle,heel,toe"', f'{col}2:{col}{last}')
dv(ws, '"Always,Start position (set-up),At the top,During the movement"', f'L2:L{last}')
dv(ws, '"Yes,No"', f'M2:M{last}')
dv(ws, '"Arc with number,Guide line,Plumb line,Floor line,Readout,Nothing"', f'N2:N{last}')
ws.freeze_panes = 'B2'

# ---------------------------------------------------------------- Faults
ws = sheet('Faults', [12, 8, 20, 44, 44, 14, 20, 12, 12, 16, 40])
cols = ['id', 'Order (1 is said first)', 'Short words on the picture', 'Spoken cue', 'Stronger words (well past the band)', 'Measurement', 'Which side of the band', 'Set-up fault?', 'A prompt, not a fault?', 'Tone', 'Cause, and what to check before it (the chain of cause)']
ex = [
    ['heelsUp', 1, 'Heels lifting', 'Keep your heels down', 'Heels down — they are coming off the floor', 'foot', 'Above', 'Yes', 'No', 'plain', 'The base. Fixes nothing else until it is right.'],
    ['toesUp', 2, 'Toes lifting', 'Keep your toes down', 'Toes down — they are coming off the floor', 'foot', 'Below', 'Yes', 'No', 'plain', ''],
    ['feetFar', 3, 'Feet too far out', 'Walk your feet in', 'Walk your feet in toward you — they are well out', 'shin', 'Above', 'Yes', 'No', 'walking in', 'Only once the foot is flat.'],
    ['feetClose', 4, 'Feet too close', 'Walk your feet out', 'Walk your feet out — your knees are out over your toes', 'shin', 'Below', 'Yes', 'No', 'walking out', ''],
    ['raise', 5, '', 'Lift your hips', '', 'hip', 'Below the raise threshold at the start', 'No', 'Yes', 'plain', 'The prompt: asked for at the start position, never red.'],
    ['hipHigh', 6, 'Hips above knees', 'Not so high — hips no higher than your knees', 'Lower your hips — they are well above your knees', 'over', 'Above', 'No', 'No', 'down', ''],
    ['hipLow', 7, 'Hips short of the line', 'Lift your hips higher', 'Higher — knees, hips and shoulders in one line', 'hip', 'Below', 'No', 'No', 'up', 'Judged at the top only.'],
]
last = table(ws, 1, cols, ex, 12, 'Order is the chain of cause: what the rest of the body stands on comes first (feet before knees before hips; hands and arms before the back before the leg). Set-up faults are coached at the start position before the movement is asked for. A prompt is what asks for the movement; it is never red. Tones: plain (one short note), up (rising two notes: lift), down (falling: lower), walking in / walking out (three-note runs), hold (two rising notes), done (three), call (the clock, twice).')
dv(ws, '"Above,Below,Outside either way,Below the raise threshold at the start"', f'G2:G{last}')
dv(ws, '"Yes,No"', f'H2:H{last}'); dv(ws, '"Yes,No"', f'I2:I{last}')
dv(ws, '"plain,up,down,walking in,walking out,hold,done,call"', f'J2:J{last}')
ws.freeze_panes = 'B2'

# ---------------------------------------------------------------- Words & sounds
ws = sheet('Words & sounds', [30, 56, 56])
header_row(ws, 1, ['Moment', 'Words (yours; blank keeps the app\'s)', 'The app\'s words today'])
words = [
    ('Into position (a hold, or the top of a rep)', 'That is it — hold'),
    ('The hold at the top is done (reps)', 'Lower slowly'),
    ('The rep counted', '1, 2, 3 … (the number alone)'),
    ('Counted, but lowered too fast', '3 — slower on the way down'),
    ('Back down before the hold was done', 'Hold it at the top next time'),
    ('The last rep', '10 reps — done'),
    ('A hold\'s target reached', '60 seconds — done'),
    ('Time called (holds)', '30 seconds left'),
    ('Nobody in the frame (every 15 s)', 'I can\'t see you — step into the camera, side on'),
    ('Next set', 'Set 2 of 3. When you are ready.'),
    ('Next set, other side', 'Set 2 of 3 — the other leg. When you are ready.'),
    ('Session over', 'All done. 3 sets, 30 reps.'),
]
for i, (k, v) in enumerate(words, 2):
    ws.cell(row=i, column=1, value=k).font = bold
    c = ws.cell(row=i, column=2); c.fill = fill_in; c.border = box; c.font = ink
    e = ws.cell(row=i, column=3, value=v); e.font = blue
r = len(words) + 4
header_row(ws, r, ['Timing rule', 'Your value', 'Default'])
timing = [('Persist: how long a fault holds before it is said (s)', 0.5), ('Cooldown: the same cue not again inside (s)', 4), ('Gap: no two cues inside (s)', 1.5), ('Settle: in position this long before the clock starts (s)', 0.7), ('Set-up wait (s)', 2), ('Can\'t-see-you interval (s)', 15), ('Deep: degrees past the band at which the stronger words are used', 18), ('Smoothing (1 = none)', 0.35)]
for i, (k, v) in enumerate(timing, r + 1):
    ws.cell(row=i, column=1, value=k).font = ink
    c = ws.cell(row=i, column=2); c.fill = fill_in; c.border = box; c.font = ink
    ws.cell(row=i, column=3, value=v).font = blue

# ---------------------------------------------------------------- Muscles
ws = sheet('Muscles', [16, 14, 16, 40])
cols = ['Muscle', 'Effort (0 to 1)', 'Role', 'Notes']
ex = [['glute', 1, 'Primary', 'The lift.'], ['ham', 0.7, 'Primary', 'Hip extension with the glutes.'], ['abs', 0.4, 'Stabiliser', 'Keeps the ribs down; stops the back arching.'], ['back', 0.3, 'Stabiliser', 'Lower back holds the line at the top.'], ['thigh', 0.2, 'Secondary', '']]
last = table(ws, 1, cols, ex, 8, 'The muscles the figure can show: shoulder, arm, forearm, thigh, ham, calf, chest, back, abs, oblique, neck, glute. Effort drives how warm the muscle is drawn.')
dv(ws, '"shoulder,arm,forearm,thigh,ham,calf,chest,back,abs,oblique,neck,glute"', f'A2:A{last}')
dv(ws, '"Primary,Secondary,Stabiliser"', f'C2:C{last}')

# ---------------------------------------------------------------- Figure
ws = sheet('Figure', [26, 16, 16, 60])
header_row(ws, 1, ['Joint angle (degrees)', 'A — start', 'B — end (blank for a hold)', 'How to read it'])
fig = [
    ('face', 'right', '', 'Which way the figure faces on screen: right or left.'),
    ('torso (from vertical, + leaning forward)', 90, 90, 'Lying flat is 90.'),
    ('neck', 0, 0, 'Head on from the torso.'),
    ('thigh (from straight down, + forward)', 180, 150, 'Lying: the thigh points along the floor.'),
    ('shin', 90, 60, ''),
    ('foot (from horizontal, + toes up)', 0, 0, ''),
    ('upper arm', 170, 170, 'By the side, on the floor.'),
    ('forearm', 170, 170, ''),
    ('far thigh / shin / arms (if different)', '', '', 'Leave blank when the far limb mirrors the near one.'),
    ('wall', 'none', '', 'behind, ahead, or none.'),
    ('hold', 'no', '', 'yes: the figure holds A. no: it moves A to B and back.'),
]
for i, (k, a, b, how) in enumerate(fig, 2):
    ws.cell(row=i, column=1, value=k).font = bold
    for col, v in ((2, a), (3, b)):
        c = ws.cell(row=i, column=col, value=v); c.font = blue; c.fill = fill_ex; c.border = box
    ws.cell(row=i, column=4, value=how).font = ink
r = len(fig) + 4
c = ws.cell(row=r, column=1, value='These are the example\'s values. Overwrite them with yours; the Animation tab of the Review page draws the figure from the same numbers and lets you drag the joints instead.'); c.font = muted; c.alignment = wrap
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

# ---------------------------------------------------------------- Checklist
ws = sheet('Checklist', [70, 14, 40])
header_row(ws, 1, ['Before the exercise is switched on', 'Done?', 'Notes'])
items = [
    'Every measurement above can be read from the chosen view (nothing hidden behind the body).',
    'Three clean takes recorded on a phone, from the phone position above.',
    'Two takes of each fault recorded, tagged on the Review page.',
    'The tuning rule passes on the Review page: quiet on every clean take, fires on every take of its fault.',
    'The starting-position rule holds on the clean takes for the whole set-up wait, and fails when the person is not in it.',
    'The raise and down thresholds cut every clean take into the right number of reps (Rep by rep on the Review page).',
    'The start text is under 120 characters and says where the phone goes and what position to take.',
    'A demo film rendered from a clean take, watched with the sound on.',
    'The numbers copied into the move\'s defaults; the About text written from the Exercise sheet.',
]
for i, it in enumerate(items, 2):
    ws.cell(row=i, column=1, value=it).font = ink; ws.cell(row=i, column=1).alignment = wrap
    c = ws.cell(row=i, column=2); c.fill = fill_in; c.border = box; c.font = ink
    ws.cell(row=i, column=3).fill = fill_in; ws.cell(row=i, column=3).border = box
dv(ws, '"Yes,No"', f'B2:B{len(items) + 1}')
r = len(items) + 3
c = ws.cell(row=r, column=1, value='When every row reads Yes, the move is ready to switch on: copy the numbers into its defaults in public/js/moves.js and write its About text from the Exercise sheet.'); c.font = muted; c.alignment = wrap
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

wb.save('docs/exercise-template.xlsx')
print('wrote docs/exercise-template.xlsx')
